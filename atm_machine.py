from openpyxl import Workbook
from datetime import datetime
import os

file_name = "/storage/emulated/0/Download/ATM_Transactions.xlsx"

balance = 5000

if os.path.exists(file_name):
    from openpyxl import load_workbook
    wb = load_workbook(file_name)
    ws = wb.active
    # Restore balance from the last recorded transaction
    if ws.max_row > 1:
        balance = ws.cell(ws.max_row, 4).value or 5000
else:
    wb = Workbook()
    ws = wb.active
    ws.title = "ATM Transactions"
    ws.append([
        "Date and Time", "Transaction",
        "Amount", "Balance", "Status"
    ])

print("\n===== ATM MACHINE =====")
print("1. Check Balance")
print("2. Deposit Money")
print("3. Withdraw Money")
print("4. Exit")

try:
    choice = int(input("Enter your choice: "))
    amount = 0
    status = ""

    if choice == 1:
        transaction = "Balance Check"
        status = "Successful"
        print("Your balance: Rs.", balance)

    elif choice == 2:
        transaction = "Deposit"
        amount = int(input("Enter deposit amount: "))

        if amount > 0:
            balance += amount
            status = "Successful"
            print("Deposit successful!")
            print("New balance: Rs.", balance)
        else:
            status = "Invalid Amount"
            print("Invalid amount!")

    elif choice == 3:
        transaction = "Withdrawal"
        amount = int(input("Enter withdrawal amount: "))

        if amount <= 0:
            status = "Invalid Amount"
            print("Invalid amount!")
        elif amount <= balance:
            balance -= amount
            status = "Successful"
            print("Please collect your cash.")
            print("Remaining balance: Rs.", balance)
        else:
            status = "Insufficient Balance"
            print("Insufficient balance!")

    elif choice == 4:
        transaction = "Exit"
        status = "Program Closed"
        print("Thank you for using ATM!")

    else:
        transaction = "Invalid Choice"
        status = "Failed"
        print("Invalid choice!")

    ws.append([
        datetime.now().strftime("%d-%m-%Y %H:%M:%S"),
        transaction, amount, balance, status
    ])

    wb.save(file_name)
    print("Excel file saved at:")
    print(file_name)

except ValueError:
    print("Please enter a valid number!")